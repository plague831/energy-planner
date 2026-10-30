"""Збирає розділи пояснювальної записки в один .docx на стилях шаблону кафедри.

Використання:
    python3 build.py <шаблон.dotx> <джерела.md> <вихід.docx> <розділ-1.md> [<розділ-2.md> ...]

Посилання в тексті пишуться як [@ключ] або [@ключ1; @ключ2] і нумеруються
в порядку першої появи через усі розділи. Перелік джерел (лише ті, на які є
посилання) додається в кінець файлу. Скрипт також перевіряє, що кожен рисунок
і кожна таблиця мають посилання в тексті, і показує джерела, на які посилань немає.

Рисунок: рядок ![Рисунок 2.1 – Назва](figures/ris-2-1.png), шлях відносно файлу розділу.
"""
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from PIL import Image, ImageFont

try:
    FONT = ImageFont.truetype('LiberationSerif-Regular.ttf', 100)
except OSError:
    FONT = None

BODY = '0'          # «0» у шаблоні: основний текст з відступом 1,25 см
H1, H2 = '1', '2'
TABLE = 'a3'        # Table Grid
TABLE_PT = 12       # один розмір шрифту в усіх таблицях роботи (вимога керівника)
FIGURE_MM = 165
CELL_MARGIN_MM = 1.0
FIGURE_MAX_H_MM = 200


def open_template(path):
    """python-docx не відкриває .dotx, тому міняємо тип головної частини."""
    tmp = Path(tempfile.mkdtemp()) / 'template.docx'
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = data.replace(
                    b'application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml',
                    b'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml')
            dst.writestr(item, data)
    doc = Document(str(tmp))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)
    return doc


def load_sources(path):
    sources = {}
    for line in Path(path).read_text(encoding='utf8').splitlines():
        m = re.match(r'- (\w+): (.+)', line)
        if m:
            sources[m.group(1)] = m.group(2)
    return sources


def number_citations(text, order):
    def repl(m):
        nums = []
        for key in re.findall(r'@(\w+)', m.group(1)):
            if key not in order:
                order.append(key)
            nums.append(order.index(key) + 1)
        return '[' + compress(sorted(set(nums))) + ']'
    return re.sub(r'\[(@[^\]]+)\]', repl, text)


def compress(nums):
    """[3, 4, 5, 8] -> «3–5, 8»: три і більше номерів поспіль подаються діапазоном."""
    parts, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        if j - i >= 2:
            parts.append(f'{nums[i]}–{nums[j]}')
        else:
            parts.extend(str(n) for n in nums[i:j + 1])
        i = j + 1
    return ', '.join(parts)


INDEX = re.compile(r'([_^])(\{[^}]+\}|[A-Za-zα0-9]+)')


def add_runs(par, text, size=None, subscripts=True):
    """Підтримує *курсив*, нижні індекси CVaR_α, P_{ch}, верхні 10^{7} і `буквальний текст` (назви полів бази)."""
    for k, chunk in enumerate(text.split('`')):
        if k % 2:
            run = par.add_run(chunk)
            if size:
                run.font.size = Pt(size)
        else:
            add_formatted(par, chunk, size, subscripts)


def add_formatted(par, text, size, subscripts):
    for i, part in enumerate(re.split(r'\*([^*]+)\*', text)):
        if not part:
            continue
        pieces, pos = [], 0
        for m in (INDEX.finditer(part) if subscripts else []):
            if m.group(1) == '^' and not m.group(2).startswith('{'):
                continue                      # верхній індекс лише у фігурних дужках
            pieces.append((part[pos:m.start()], None))
            pieces.append((m.group(2).strip('{}'), m.group(1)))
            pos = m.end()
        pieces.append((part[pos:], None))
        for piece, kind in pieces:
            if not piece:
                continue
            run = par.add_run(piece)
            run.italic = bool(i % 2) or None
            if kind == '_':
                run.font.subscript = True
            elif kind == '^':
                run.font.superscript = True
            if size:
                run.font.size = Pt(size)


def heading(doc, text, style, page_break=False):
    p = doc.add_paragraph(style=style)
    if page_break:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0 if page_break or style == H1 else 21)
    p.paragraph_format.space_after = Pt(21)
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def caption_above(doc, text):
    p = doc.add_paragraph(style=BODY)
    p.paragraph_format.first_line_indent = 0
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, text)


def table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = doc.styles['Table Grid']
    t.autofit = False
    # Ширина колонки пропорційна середній довжині тексту в ній, але не менша за найдовше слово,
    # щоб слова не розривалися; разом 170 мм (ширина тексту на сторінці).
    weights, minimum = [], []
    for c in range(len(rows[0])):
        texts = [re.sub(r'[`*{}^]', '', row[c]) for row in rows]
        longest = max((w for x in texts for w in x.split()), key=word_mm, default='')
        weights.append(max(sum(map(len, texts)) / len(texts), 4))
        minimum.append(word_mm(longest) + 2 * CELL_MARGIN_MM + 0.6)
    if sum(minimum) > 170:
        minimum = [m * 170 / sum(minimum) for m in minimum]
    widths = [170 * w / sum(weights) for w in weights]
    for _ in range(len(widths)):
        narrow = [c for c, w in enumerate(widths) if w < minimum[c] - 1e-6]
        if not narrow:
            break
        fixed = set(c for c, w in enumerate(widths) if w <= minimum[c] + 1e-6) | set(narrow)
        rest = [c for c in range(len(widths)) if c not in fixed]
        for c in fixed:
            widths[c] = minimum[c]
        free = 170 - sum(widths[c] for c in fixed)
        for c in rest:
            widths[c] = free * weights[c] / sum(weights[k] for k in rest)
    margins(t)
    for c, col in enumerate(t.columns):
        col.width = Mm(widths[c])
        for cell in col.cells:
            cell.width = col.width
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            p = t.cell(r, c).paragraphs[0]
            p.style = doc.styles['Normal']
            p.paragraph_format.line_spacing = 1.0
            if r == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # невелика таблиця не розривається між сторінками: рядки «тримаються» за наступні
            p.paragraph_format.keep_with_next = r < len(rows) - 1
            add_runs(p, cell, size=TABLE_PT)
            if r == 0:
                for run in p.runs:
                    run.bold = False
        cant_split = OxmlElement('w:cantSplit')
        t.rows[r]._tr.get_or_add_trPr().append(cant_split)
    # повторювати шапку на кожній сторінці
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement('w:tblHeader')
    hdr.set(qn('w:val'), 'true')
    tr_pr.append(hdr)
    doc.add_paragraph(style=BODY)


def word_mm(word):
    """Ширина слова шрифтом таблиці в мм (Liberation Serif метрично збігається з Times New Roman)."""
    if FONT is None:
        return len(word) * 2.1
    return FONT.getlength(word) / 100 * TABLE_PT * 25.4 / 72


def margins(t):
    """Поля комірок ліворуч і праворуч по CELL_MARGIN_MM замість стандартних 1,9 мм."""
    mar = OxmlElement('w:tblCellMar')
    for side in ('left', 'right'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(int(CELL_MARGIN_MM / 25.4 * 1440)))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    t._tbl.tblPr.append(mar)


def formula(doc, text, number):
    p = doc.add_paragraph(style=BODY)
    p.paragraph_format.first_line_indent = 0
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Pt(240), alignment=1)   # центр
    pf.tab_stops.add_tab_stop(Pt(482), alignment=2)   # праворуч
    p.add_run('\t')
    add_runs(p, text)
    p.add_run('\t' + number)


def figure(doc, image, caption):
    p = doc.add_paragraph(style=BODY)
    p.paragraph_format.first_line_indent = 0
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # широкий рисунок займає 165 мм, високий обмежено 200 мм заввишки, щоб він уміщався з підписом
    with Image.open(image) as im:
        w, h = im.size
    p.add_run().add_picture(str(image), width=Mm(min(FIGURE_MM, FIGURE_MAX_H_MM * w / h)))
    c = doc.add_paragraph(style=BODY)
    c.paragraph_format.first_line_indent = 0
    c.paragraph_format.space_after = Pt(14)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(c, caption)


def build(template, sources_path, out, chapters):
    doc = open_template(template)
    sources = load_sources(sources_path)
    order = []
    captions, text = [], []
    first = True
    for chapter in chapters:
        base = Path(chapter).parent
        lines = Path(chapter).read_text(encoding='utf8').splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()
            if not line:
                i += 1
                continue
            m = re.fullmatch(r'!\[(.+)\]\((.+)\)', line)
            if line.startswith('# '):
                heading(doc, line[2:], H1, page_break=not first)
                first = False
            elif line.startswith('## '):
                heading(doc, line[3:], H2)
            elif m:
                figure(doc, base / m.group(2), m.group(1))
                captions.append(m.group(1))
            elif line.startswith('Таблиця '):
                caption_above(doc, line)
                captions.append(line)
            elif line.startswith('|'):
                rows = []
                while i < len(lines) and lines[i].startswith('|'):
                    cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                    if not all(re.fullmatch(r'-+', c) for c in cells):
                        rows.append([number_citations(c, order) for c in cells])
                    i += 1
                table(doc, rows)
                continue
            elif re.search(r'[,.] \((\d+\.\d+)\)$', line):
                m = re.search(r'^(.*[,.]) \((\d+\.\d+)\)$', line)
                formula(doc, m.group(1), f'({m.group(2)})')
            else:
                p = doc.add_paragraph(style=BODY)
                add_runs(p, number_citations(line, order))
                text.append(line)
            i += 1

    heading(doc, 'ПЕРЕЛІК ДЖЕРЕЛ ПОСИЛАННЯ', H1, page_break=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    for n, key in enumerate(order, 1):
        if key not in sources:
            sys.exit(f'Немає опису джерела для ключа {key}')
        p = doc.add_paragraph(style=BODY)
        add_runs(p, f'{n}. {sources[key]}', subscripts=False)
    doc.save(out)
    print(f'{out}: джерел {len(order)}')
    check(captions, '\n'.join(text), sources, order)


def check(captions, text, sources, order):
    """Попередження про рисунки й таблиці без посилань у тексті і про невикористані джерела."""
    for cap in captions:
        kind, num = re.match(r'(Рисунок|Таблиця) (\d+\.\d+)', cap).groups()
        ref = r'рис\. ' if kind == 'Рисунок' else r'табл(\.|иц\w+) '
        if not re.search(ref + r'(\d+\.\d+, )*' + re.escape(num) + r'\b', text):
            print(f'Попередження: {kind.lower()} {num} не згадується в тексті')
    unused = [k for k in sources if k not in order]
    if unused:
        print('Джерела без посилань (у перелік не потрапили):', ', '.join(unused))


if __name__ == '__main__':
    if len(sys.argv) < 5:
        sys.exit(__doc__)
    build(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4:])
